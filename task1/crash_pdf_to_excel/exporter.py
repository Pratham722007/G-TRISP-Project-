import pandas as pd
from pathlib import Path
from typing import List, Dict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class CrashExporter:
    """
    Logic for writing extracted crash data to a structured Excel file.
    """

    def __init__(self, output_path: Path):
        self.output_path = output_path
        # Define header color and font
        self.header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True)
        self.center_alignment = Alignment(horizontal="center", vertical="center")

    def export(self, data_list: List[Dict[str, str]]):
        """
        Writes a list of dictionaries to an Excel file.
        Each dictionary represents one row.
        """
        if not data_list:
            print("  Warning: No data to export.")
            return

        # Ensure output directory exists
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        # Create workbook and sheet
        wb = Workbook()
        ws = wb.active
        ws.title = "Crash_Data"

        # Get headers from the first dictionary keys
        headers = list(data_list[0].keys())

        # Write header row
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.center_alignment

        # Write data rows
        for row_num, entry in enumerate(data_list, 2):
            for col_num, header in enumerate(headers, 1):
                ws.cell(row=row_num, column=col_num, value=entry.get(header, ""))

        # Styling: Auto-size columns (cap at 40)
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 40)
            ws.column_dimensions[column].width = adjusted_width

        # Freeze the top header row
        ws.freeze_panes = "A2"

        # Save the workbook
        wb.save(self.output_path)
        print(f"  Excel file saved to: {self.output_path}")
